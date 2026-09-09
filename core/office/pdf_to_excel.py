"""Best-effort local PDF to Excel (XLSX) export with table and column detection."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Callable

import fitz
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font

from core.utils.file_utils import atomic_output, ensure_distinct_paths
from core.utils.validation import validate_pdf

Progress = Callable[[int, str], None]

WRAP = Alignment(wrap_text=True, vertical="top")
COLUMN_WIDTH = 100.0
ROW_HEIGHT_PX = 18.0
CELL_GAP = 10.0


def _inside_any(px: float, py: float, rects: list[tuple[float, float, float, float]]) -> bool:
    return any(r[0] <= px <= r[2] and r[1] <= py <= r[3] for r in rects)


def _detect_tables(page):
    for strategy in ("lines", "text"):
        try:
            found = (page.find_tables(strategy=strategy) or []).tables
            if found: return found
        except Exception:
            continue
    return []


def _split_cells(words: list[tuple[float, float, float, float, str]]) -> list[str]:
    cells: list[str] = []
    current: list[str] = []
    previous_x1: float | None = None
    for x0, _y0, x1, _y1, text in sorted(words, key=lambda item: (item[0], item[2])):
        if current and previous_x1 is not None and x0 - previous_x1 > CELL_GAP:
            cells.append(" ".join(current)); current = []
        current.append(text); previous_x1 = x1
    if current: cells.append(" ".join(current))
    return cells or [""]


def _page_bands(page) -> list[tuple[float, str, object, float, float]]:
    tables = _detect_tables(page)
    regions = [table.bbox for table in tables]
    bands: list[tuple[float, str, object, float, float]] = []
    groups: dict[tuple[int, int], list] = {}
    for word in page.get_text("words"):
        x0, y0, x1, y1, text = word[0], word[1], word[2], word[3], word[4]
        if not text.strip(): continue
        groups.setdefault((word[5], word[6]), []).append((x0, y0, x1, y1, text))
    for (block, _line), words in groups.items():
        y0 = min(item[1] for item in words); x0 = min(item[0] for item in words); y1 = max(item[3] for item in words); x1 = max(item[2] for item in words)
        if _inside_any((x0 + x1) / 2, (y0 + y1) / 2, regions): continue
        bands.append((y0, "text", _split_cells(words), x0, y1))
    for table in tables:
        x0, y0, x1, y1 = table.bbox
        rows = table.extract()
        bands.append((y0, "table", rows, x0, y1))
    for info in page.get_image_info(xrefs=True):
        x0, y0, x1, y1 = info["bbox"]
        if _inside_any((x0 + x1) / 2, (y0 + y1) / 2, regions): continue
        if info.get("xref") is None: continue
        bands.append((y0, "image", (info["xref"], info), x0, y1))
    return sorted(bands, key=lambda band: (band[0], band[3]))


def pdf_to_excel(
    source: str | Path, destination: str | Path, *, progress: Progress | None = None,
) -> Path:
    """Export each page as a worksheet: tables as grids, text as rows, images embedded."""
    info = validate_pdf(source)
    output = Path(destination).resolve()
    ensure_distinct_paths(info.path, output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("PDF to Excel output must use .xlsx.")
    workbook = Workbook(); workbook.remove(workbook.active)
    created = 0
    with fitz.open(info.path) as pdf:
        for page_index, page in enumerate(pdf):
            bands = _page_bands(page)
            if not bands:
                continue
            created += 1
            sheet = workbook.create_sheet(title=f"Page {page_index + 1}"[:31])
            sheet.column_dimensions["A"].width = COLUMN_WIDTH
            row = 1
            for _y0, kind, payload, _x0, _y1 in bands:
                if kind == "text":
                    for column, text in enumerate(payload):
                        cell = sheet.cell(row=row, column=column + 1, value=text)
                        cell.alignment = WRAP
                        sheet.column_dimensions[_col_index(column)].width = COLUMN_WIDTH
                    row += 1
                elif kind == "table":
                    for table_row in payload:
                        if not any(value is not None and str(value).strip() for value in table_row): continue
                        for column, value in enumerate(table_row):
                            if value is None: continue
                            cell = sheet.cell(row=row, column=column + 1, value=str(value))
                            cell.alignment = WRAP
                            sheet.column_dimensions[_col_index(column)].width = min(COLUMN_WIDTH, max(sheet.column_dimensions[_col_index(column)].width or COLUMN_WIDTH, len(str(value)) * 1.35 + 3))
                        row += 1
                    row += 1
                elif kind == "image":
                    xref, info_dict = payload
                    try:
                        data = pdf.extract_image(xref)["image"]
                        canvas = XlImage(BytesIO(data))
                        canvas.width = int(info_dict["width"]); canvas.height = int(info_dict["height"])
                        sheet.add_image(canvas, f"A{row}")
                        sheet.row_dimensions[row].height = info_dict["height"]
                        row += max(1, round(info_dict["height"] / ROW_HEIGHT_PX))
                    except Exception:
                        continue
            if progress: progress(round((page_index + 1) / pdf.page_count * 95), f"Exported page {page_index + 1}")
    if created == 0:
        raise ValueError("The PDF contains no extractable text, tables, or images.")
    with atomic_output(output) as temporary: workbook.save(temporary)
    if progress: progress(100, output.name)
    return output


def _col_index(column: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(column + 1)