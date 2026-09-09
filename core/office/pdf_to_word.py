"""Best-effort local PDF to DOCX conversion preserving lines, formatting, tables, and images."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Inches, Pt

from core.utils.file_utils import atomic_output, ensure_distinct_paths
from core.utils.validation import validate_pdf

MAX_IMAGE_WIDTH_IN = 6.5
BOLD_FLAG = 16
ITALIC_FLAG = 2


def _inside_any(px: float, py: float, rects: list[tuple[float, float, float, float]]) -> bool:
    return any(r[0] <= px <= r[2] and r[1] <= py <= r[3] for r in rects)


def _detect_tables(page):
    for strategy in ("lines", "text"):
        try:
            found = (page.find_tables(strategy=strategy) or []).tables
            if found: return found
        except Exception:
            continue
    return []


def _page_bands(page):
    tables = _detect_tables(page)
    regions = [table.bbox for table in tables]
    bands: list[tuple[float, str, object, float, float]] = []
    data = page.get_text("dict")
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            bbox = line["bbox"]
            if _inside_any((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2, regions): continue
            spans = []
            for span in line["spans"]:
                text = span.get("text")
                if not text or not text.strip(): continue
                flags = int(span.get("flags") or 0)
                spans.append((text, float(span.get("size") or 11), bool(flags & BOLD_FLAG), bool(flags & ITALIC_FLAG)))
            if spans: bands.append((bbox[1], "text", spans, bbox[0], bbox[3]))
    for table in tables:
        x0, y0, x1, y1 = table.bbox
        rows = table.extract()
        if rows: bands.append((y0, "table", rows, x0, y1))
    for info in page.get_image_info(xrefs=True):
        x0, y0, x1, y1 = info["bbox"]
        if _inside_any((x0 + x1) / 2, (y0 + y1) / 2, regions): continue
        if info.get("xref") is None: continue
        bands.append((y0, "image", (info["xref"], info), x0, y1))
    return sorted(bands, key=lambda band: (band[0], band[3]))


def _add_table(document: Document, rows) -> None:
    cleaned = [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]
    if not cleaned: return
    columns = max((len(row) for row in cleaned), default=1) or 1
    table = document.add_table(rows=len(cleaned), cols=columns)
    table.style = "Table Grid"
    for row_index, row in enumerate(cleaned):
        for column, value in enumerate(row):
            if value is None: continue
            table.cell(row_index, column).text = str(value)
    document.add_paragraph()


def pdf_to_word(source: str | Path, destination: str | Path, *, include_images: bool = True, progress=None) -> Path:
    info = validate_pdf(source)
    output = Path(destination).resolve()
    ensure_distinct_paths(info.path, output)
    if output.suffix.lower() != ".docx": raise ValueError("PDF to Word output must use .docx.")
    word = Document(); word.core_properties.title = info.path.stem
    total = 0
    with fitz.open(info.path) as pdf:
        for page_index, page in enumerate(pdf):
            total += 1
            bands = _page_bands(page)
            width = page.rect.width
            for _y0, kind, payload, _x0, _y1 in bands:
                if kind == "text":
                    paragraph = word.add_paragraph()
                    for text, size, bold, italic in payload:
                        run = paragraph.add_run(text)
                        run.bold = bool(bold); run.italic = bool(italic)
                        if size: run.font.size = Pt(size)
                elif kind == "table":
                    _add_table(word, payload)
                elif kind == "image" and include_images:
                    xref, info_dict = payload
                    try:
                        extracted = pdf.extract_image(xref)
                        stream = BytesIO(extracted["image"])
                        image_width = float(info_dict["width"])
                        if width > 0:
                            size_in = min(MAX_IMAGE_WIDTH_IN, max(1.0, MAX_IMAGE_WIDTH_IN * image_width / width))
                        else:
                            size_in = MAX_IMAGE_WIDTH_IN
                        word.add_picture(stream, width=Inches(size_in))
                    except Exception:
                        continue
            if page_index < pdf.page_count - 1: word.add_page_break()
            if progress: progress(round((page_index + 1) / pdf.page_count * 95), f"Converting page {page_index + 1} of {pdf.page_count}")
    if total == 0:
        raise ValueError("The PDF contains no pages.")
    with atomic_output(output) as temporary: word.save(temporary)
    if progress: progress(100, output.name)
    return output