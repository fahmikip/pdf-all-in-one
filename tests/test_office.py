from pathlib import Path

import pymupdf
import pytest
from core.office.libreoffice_converter import LibreOfficeUnavailableError, find_libreoffice, office_to_pdf
from core.office.pdf_to_excel import pdf_to_excel
from core.office.pdf_to_word import pdf_to_word
from docx import Document
from openpyxl import load_workbook
from PIL import Image


def test_pdf_to_word_extracts_text(sample_pdf: Path, tmp_path: Path) -> None:
    output = pdf_to_word(sample_pdf, tmp_path / "output.docx", include_images=False)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Page 1" in text and "Page 3" in text


def _content_pdf(path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=300, height=400)
    page.insert_text((40, 60), "Item")
    page.insert_text((40, 120), "Pen")
    image = Path(path.parent) / "photo.png"
    Image.new("RGB", (80, 60), (200, 40, 40)).save(image)
    page.insert_image(pymupdf.Rect(40, 160, 120, 220), filename=str(image))
    document.save(path)
    document.close()
    return path


def test_pdf_to_excel_exports_text_and_images(tmp_path: Path) -> None:
    pdf = _content_pdf(tmp_path / "content.pdf")
    output = pdf_to_excel(pdf, tmp_path / "content.xlsx")
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Page 1"]
    sheet = workbook.active
    values = [sheet.cell(row=r, column=1).value for r in range(1, 4)]
    assert "Item" in values and "Pen" in values
    assert sheet._images, "expected an embedded image"


def test_pdf_to_excel_empty_page_raises(sample_pdf: Path, tmp_path: Path) -> None:
    blank = tmp_path / "blank.pdf"
    document = pymupdf.open()
    document.new_page(width=300, height=400)
    document.save(blank)
    document.close()
    with pytest.raises(ValueError):
        pdf_to_excel(blank, tmp_path / "out.xlsx")


def test_pdf_to_excel_rejects_non_xlsx(sample_pdf: Path, tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        pdf_to_excel(sample_pdf, tmp_path / "out.pdf")


def test_missing_libreoffice_has_friendly_error(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr("core.office.libreoffice_converter.shutil.which", lambda name: None)
    monkeypatch.setattr("core.office.libreoffice_converter.Path.is_file", lambda self: False)
    assert find_libreoffice() is None
    with pytest.raises(LibreOfficeUnavailableError):
        office_to_pdf(tmp_path / "document.docx", tmp_path / "document.pdf")


def test_find_libreoffice_uses_manual_path(tmp_path: Path) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    assert find_libreoffice(soffice) == soffice.resolve()


def test_office_to_pdf_rejects_missing_file(monkeypatch, tmp_path: Path) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    with pytest.raises(ValueError):
        office_to_pdf(tmp_path / "missing.docx", tmp_path / "out.pdf", executable=soffice)


def test_office_to_pdf_rejects_unsupported_extension(monkeypatch, tmp_path: Path) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    document = tmp_path / "archive.html"
    document.write_text("<html></html>", encoding="utf-8")
    with pytest.raises(ValueError):
        office_to_pdf(document, tmp_path / "out.pdf", executable=soffice)


def test_office_to_pdf_rejects_wrong_output_suffix(monkeypatch, tmp_path: Path) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    document = tmp_path / "document.docx"
    document.write_bytes(b"PK")
    with pytest.raises(ValueError):
        office_to_pdf(document, tmp_path / "out.docx", executable=soffice)


def test_office_to_pdf_success(tmp_path: Path, monkeypatch) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    document = tmp_path / "document.docx"
    document.write_bytes(b"PK")
    captured: list[list[str]] = []

    def fake_run(command, *, capture_output=True, text=True, timeout=600, check=False, creationflags=0):
        captured.append(list(command))
        outdir = Path(command[9])
        generated = outdir / f"{Path(command[10]).stem}.pdf"
        pdf = pymupdf.open()
        pdf.new_page(width=200, height=200)
        pdf.save(generated)
        pdf.close()
        return type("Result", (), {"returncode": 0, "stdout": "", "stderr": ""})()

    monkeypatch.setattr("core.office.libreoffice_converter.subprocess.run", fake_run)
    destination = tmp_path / "out.pdf"
    progress: list[tuple[int, str]] = []
    result = office_to_pdf(document, destination, executable=soffice, progress=lambda v, d: progress.append((v, d)))
    assert result == destination.resolve()
    with pymupdf.open(destination) as pdf:
        assert pdf.page_count == 1
    assert captured and "soffice" in captured[0][0]
    assert progress == [(10, "Starting LibreOffice"), (90, "Finalizing PDF"), (100, destination.name)]


def test_office_to_pdf_raises_on_subprocess_failure(tmp_path: Path, monkeypatch) -> None:
    soffice = tmp_path / "soffice.exe"
    soffice.write_bytes(b"MZ")
    document = tmp_path / "document.docx"
    document.write_bytes(b"PK")

    def fake_run(command, *, capture_output=True, text=True, timeout=600, check=False, creationflags=0):
        return type("Result", (), {"returncode": 1, "stdout": "oops", "stderr": ""})()

    monkeypatch.setattr("core.office.libreoffice_converter.subprocess.run", fake_run)
    with pytest.raises(RuntimeError, match="oops"):
        office_to_pdf(document, tmp_path / "out.pdf", executable=soffice)


def _table_pdf(path: Path) -> Path:
    document = pymupdf.open()
    page = document.new_page(width=400, height=400)
    page.insert_text((40, 40), "Quarterly", fontsize=14, fontname="hebo")
    header = ("Date", "Amount")
    rows = (("Jan", "100"), ("Feb", "250"), ("Mar", "40"))
    left, right = 40, 190
    top = 70
    baseline = [top]
    for r, _row in enumerate(rows):
        baseline.append(top + 30 + r * 30)
    page.draw_line((left, top), (right, top))
    for y in baseline[1:]:
        page.draw_line((left, y), (right, y))
    page.draw_line((left, top), (left, baseline[-1]))
    page.draw_line((right, top), (right, baseline[-1]))
    for column, cell in enumerate(header):
        page.insert_text((left + 25 + column * 90, baseline[0] + 4), cell, fontname="hebo")
    for r, row in enumerate(rows):
        for column, value in enumerate(row):
            page.insert_text((left + 25 + column * 90, baseline[r + 1] + 4), value, fontname="helv")
    document.save(path)
    document.close()
    return path


def test_pdf_to_word_detects_table_and_bold(sample_pdf: Path, tmp_path: Path) -> None:
    pdf = _table_pdf(tmp_path / "table.pdf")
    output = pdf_to_word(pdf, tmp_path / "table.docx", include_images=False)
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Quarterly" in text
    bold = [run.bold for paragraph in document.paragraphs for run in paragraph.runs if run.text]
    assert any(bold), "expected at least one bold run"
    assert document.tables, "expected a real Word table"
    assert document.tables[0].cell(0, 1).text == "Amount"
    assert document.tables[0].cell(1, 1).text == "100"


def test_pdf_to_excel_detects_table_grid(tmp_path: Path) -> None:
    pdf = _table_pdf(tmp_path / "table.pdf")
    output = pdf_to_excel(pdf, tmp_path / "table.xlsx")
    workbook = load_workbook(output)
    sheet = workbook.active
    assert sheet["A1"].value == "Quarterly"
    header = [sheet.cell(row=2, column=column).value for column in (1, 2)]
    assert header == ["Date", "Amount"]
    assert sheet["B4"].value == "250"
